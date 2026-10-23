"""Saídas da emissão avulsa: ZIP, CSV, envio por e-mail, modelo da planilha e relatório por turma."""

import io
from datetime import timedelta

from django.contrib import messages
from django.contrib.auth.decorators import login_required
from django.core.exceptions import PermissionDenied
from django.http import FileResponse, Http404, HttpResponse
from django.shortcuts import get_object_or_404, redirect, render
from django.urls import reverse
from django.utils import timezone
from django.utils.text import slugify
from django.views.decorators.http import require_POST

from core.csvsafe import excel_writer
from core.models import EmailSettings
from core.permissions import can_manage_batch

from .. import zips
from ..batch import (
    active,
    build_zip,
    queue_emails,
)
from ..batch_import import EVENT_COLUMNS
from ..models import ZipExport
from .common import FORBIDDEN, batch_filename, closed_redirect, get_managed_batch


@login_required
def batch_zip(request, pk):
    batch = get_managed_batch(request, pk)
    if batch is None:
        raise PermissionDenied(FORBIDDEN)
    if active(batch).count() > zips.ZIP_DIRECT_LIMIT:  # lote grande: monta em segundo plano (DES-02)
        messages.info(request, 'Este lote é grande: use "Gerar ZIP" e baixe quando ficar pronto.')
        return redirect('batch_detail', pk=batch.pk)
    return FileResponse(
        build_zip(batch, request=request), as_attachment=True, filename=batch_filename(batch, 'certificados.zip')
    )


@login_required
@require_POST
def batch_zip_request(request, pk):
    batch = get_managed_batch(request, pk)
    if batch is None:
        raise PermissionDenied(FORBIDDEN)
    zips.request_export(batch, request.user)
    messages.success(request, 'O ZIP está sendo gerado. Esta página atualiza sozinha; o link aparece aqui.')
    return redirect('batch_detail', pk=batch.pk)


@login_required
def batch_zip_file(request, pk, export_pk):
    batch = get_managed_batch(request, pk)
    if batch is None:
        raise PermissionDenied(FORBIDDEN)
    export = get_object_or_404(batch.zip_exports, pk=export_pk, status=ZipExport.Status.READY)
    if not export.file or export.created_at < timezone.now() - timedelta(hours=zips.ZIP_KEEP_HOURS):
        raise Http404
    return FileResponse(export.file.open('rb'), as_attachment=True, filename=batch_filename(batch, 'certificados.zip'))


@login_required
def batch_csv(request, pk):
    batch = get_managed_batch(request, pk)
    if batch is None:
        raise PermissionDenied(FORBIDDEN)
    response = HttpResponse(content_type='text/csv; charset=utf-8')
    response['Content-Disposition'] = f'attachment; filename="{batch_filename(batch, "certificados")}.csv"'
    response.write('﻿')  # BOM: o Excel abre com acentos corretos
    writer = excel_writer(response)
    writer.writerow(
        [
            'Nome',
            'Matrícula',
            'Turma',
            'E-mail',
            'Código',
            'Link de validação',
            'Data do certificado',
            'Gerado no sistema em',
            'Enviado por e-mail em',
            'Situação',
            'Revogado em',
            'Motivo da revogação',
        ]
    )
    for c in batch.certificates.order_by('participant_name'):
        writer.writerow([
            c.participant_name,
            c.participant_document,
            c.participant_group,
            c.participant_email,
            c.code,
            request.build_absolute_uri(reverse('certificate_verify_code', args=[c.code])),
            c.certificate_date.strftime('%d/%m/%Y'),
            timezone.localtime(c.issued_at).strftime('%d/%m/%Y %H:%M'),
            timezone.localtime(c.emailed_at).strftime('%d/%m/%Y %H:%M') if c.emailed_at else '',
            'Revogado' if c.is_revoked else 'Válido',
            timezone.localtime(c.revoked_at).strftime('%d/%m/%Y %H:%M') if c.revoked_at else '',
            c.revoked_reason,
        ])  # fmt: skip
    return response


@login_required
@require_POST
def batch_email(request, pk):
    batch = get_managed_batch(request, pk)
    if batch is None:
        raise PermissionDenied(FORBIDDEN)
    if closed := closed_redirect(request, batch):
        return closed
    if not EmailSettings.load().enabled:
        messages.error(request, 'O envio de e-mails está desligado. Configure em Configurações > E-mail.')
        return redirect('batch_detail', pk=batch.pk)
    queued, waiting = queue_emails(batch, resend=request.POST.get('resend') == '1', request=request)
    if queued:
        plural = 's' if queued != 1 else ''
        messages.success(
            request,
            f'{queued} certificado{plural} na fila de envio. Os e-mails saem em instantes, um por um; '
            'atualize esta página para acompanhar.',
        )
    elif waiting:
        messages.info(
            request,
            f'{waiting} e-mail{"s" if waiting != 1 else ""} deste lote já estava{"m" if waiting != 1 else ""} na fila.',
        )
    else:
        messages.info(request, 'Não havia certificados pendentes de envio.')
    return redirect('batch_detail', pk=batch.pk)


@login_required
def batch_template(request):
    """Planilha modelo, já com o cabeçalho que a importação reconhece. ?modelo=completo: com as colunas do evento."""
    if not can_manage_batch(request.user):
        raise PermissionDenied(FORBIDDEN)
    from openpyxl import Workbook
    from openpyxl.styles import Font

    full = request.GET.get('modelo') == 'completo'
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = 'Participantes'
    event_headers = list(EVENT_COLUMNS.values()) if full else []
    headers = event_headers + ['Nome', 'Matrícula', 'Turma', 'E-mail']
    sheet.append(headers)
    widths = [36, 10, 13, 13, 14, 20][: len(event_headers)] + [40, 16, 18, 36]
    for index, (cell, width) in enumerate(zip(sheet[1], widths, strict=True), start=1):
        cell.font = Font(bold=True)
        sheet.column_dimensions[cell.column_letter].width = width
        for row in range(2, 501):
            if headers[index - 1] == 'Matrícula':
                sheet.cell(row=row, column=index).number_format = '@'  # texto: não perde zeros à esquerda
            elif headers[index - 1] in ('Início', 'Término', 'Data do certificado'):
                sheet.cell(row=row, column=index).number_format = 'DD/MM/YYYY'
    if full:
        guide = workbook.create_sheet('Como preencher')
        guide.column_dimensions['A'].width = 24
        guide.column_dimensions['B'].width = 90
        for row in [
            ('Coluna', 'Como preencher'),
            ('Evento', 'Obrigatória. Título como sai no certificado, em negrito e entre aspas.'),
            ('Tipo', 'Curso, Evento ou Agenda. Em branco: Evento.'),
            ('Início', 'Obrigatória. Data no formato dd/mm/aaaa.'),
            ('Término', 'Em branco: o mesmo dia do início.'),
            ('Carga horária', 'Só o número de horas (ex.: 20). Em branco: o certificado não cita a duração.'),
            ('Data do certificado', 'A data impressa no certificado. Em branco: a data de término.'),
            ('Nome', 'Obrigatória. Um participante por linha.'),
            ('Matrícula, Turma, E-mail', 'Opcionais. Com e-mail, dá para enviar o certificado por e-mail.'),
            ('', ''),
            ('Um evento por planilha', 'As colunas do evento podem ser preenchidas só na primeira linha. Se uma '
             'linha trouxer um evento diferente, a planilha é recusada.'),
            ('Esta aba', 'É só de instruções e não é lida. Os participantes vão na aba "Participantes".'),
        ]:  # fmt: skip
            guide.append(row)
        for cell in guide[1]:
            cell.font = Font(bold=True)
    buffer = io.BytesIO()
    workbook.save(buffer)
    response = HttpResponse(
        buffer.getvalue(), content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    )
    filename = 'modelo-planilha-completa.xlsx' if full else 'modelo-participantes.xlsx'
    response['Content-Disposition'] = f'attachment; filename="{filename}"'
    return response


@login_required
def group_report(request):
    """Relatório de horas por turma (certificates/group_report.py), na tela ou em CSV."""
    from .. import group_report as reports

    if not can_manage_batch(request.user):
        raise PermissionDenied(FORBIDDEN)
    groups = reports.groups_for(request.user)
    group = request.GET.get('turma', '')
    semester = request.GET.get('semestre', '')
    rows, semesters = reports.report(request.user, group, semester) if group in groups else ([], [])
    if group in groups and request.GET.get('formato') == 'csv':
        response = HttpResponse(content_type='text/csv; charset=utf-8')
        name = slugify(f'horas-{group}-{semester}'.strip('-')) or 'horas'
        response['Content-Disposition'] = f'attachment; filename="{name}.csv"'
        response.write('\ufeff')  # o Excel reconhece acentos
        writer = excel_writer(response)
        writer.writerow(['Participante', 'Matrícula/CPF', 'Certificados', 'Horas', 'Eventos'])
        for person, document, items, hours in rows:
            writer.writerow([person, document, len(items), hours, '; '.join(c.full_title for c in items)])
        return response
    return render(
        request,
        'certificates/group_report.html',
        {
            'groups': groups,
            'group': group,
            'semester': semester,
            'semesters': semesters,
            'rows': rows,
            'total': sum(hours for *_, hours in rows),
        },
    )
