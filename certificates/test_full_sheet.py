"""Emissão avulsa com a planilha completa: dados do evento e participantes no mesmo arquivo."""

import io
import shutil
import tempfile
from datetime import date
from pathlib import Path

from django.test import TestCase, override_settings
from django.urls import reverse
from openpyxl import load_workbook

from accounts.models import User
from activities.test_utils import make_user

from .batch_import import SpreadsheetError, column_kind, read_event
from .models import CertificateBatch
from .test_batch import csv_file, xlsx

HEADER = ('Evento', 'Tipo', 'Início', 'Término', 'Carga horária', 'Data do certificado', 'Nome', 'Matrícula', 'E-mail')


class ReadEventTests(TestCase):
    def test_colunas_reconhecidas(self):
        expected = {
            'Evento': 'title', 'Nome do evento': 'title', 'Título': 'title', 'Tipo': 'activity_type',
            'Início': 'start_date', 'Data de início': 'start_date', 'Data': 'start_date', 'Término': 'end_date',
            'Data final': 'end_date', 'Carga horária': 'workload_hours', 'Data do certificado': 'issue_date',
            'Data de emissão': 'issue_date', 'Nome': 'name', 'Data de nascimento': None,
        }  # fmt: skip
        self.assertEqual({h: column_kind(h) for h in expected}, expected)

    def test_evento_so_na_primeira_linha(self):
        table = [
            HEADER,
            ('Semana de TI', 'curso', date(2026, 9, 1), '03/09/2026', '20 horas', None, 'Ana', '1', ''),
            (None, None, None, None, None, None, 'Bia', '2', ''),
            ('Semana de TI', '', '01/09/2026', '', 20, '', 'Caio', '3', ''),  # repetido igual: tudo bem
        ]
        event = read_event(table)
        self.assertEqual(
            event,
            {'title': 'Semana de TI', 'activity_type': 'COURSE', 'start_date': date(2026, 9, 1),
             'end_date': date(2026, 9, 3), 'workload_hours': 20, 'issue_date': date(2026, 9, 3)},
        )  # fmt: skip

    def test_padroes(self):
        event = read_event([('Evento', 'Início', 'Nome'), ('Palestra', '10/08/2026', 'Ana')])
        self.assertEqual((event['activity_type'], event['end_date'], event['issue_date']), ('EVENT', date(2026, 8, 10),
                         date(2026, 8, 10)))  # fmt: skip
        self.assertIsNone(event['workload_hours'])

    def test_um_evento_por_planilha(self):
        table = [
            HEADER,
            ('Semana de TI', '', '01/09/2026', '', 20, '', 'Ana', '1', ''),
            ('Semana de TI', '', '01/09/2026', '', 8, '', 'Bia', '2', ''),
        ]
        with self.assertRaisesMessage(SpreadsheetError, 'A linha 3 traz Carga horária "8", diferente da linha 2'):
            read_event(table)
        table[2] = ('Outra palestra', '', '01/09/2026', '', '', '', 'Bia', '2', '')
        with self.assertRaisesMessage(SpreadsheetError, 'um evento só'):
            read_event(table)

    def test_valores_que_nao_dao_para_entender(self):
        for cells, message in (
            (('Semana', '', '31/02/2026', '', '', '', 'Ana', '', ''), 'coluna Início: "31/02/2026"'),
            (('Semana', 'Palestra', '01/09/2026', '', '', '', 'Ana', '', ''), 'Use Curso, Evento ou Agenda'),
            (('Semana', '', '01/09/2026', '', 'vinte', '', 'Ana', '', ''), 'Use só o número de horas'),
        ):
            with self.assertRaisesMessage(SpreadsheetError, message):
                read_event([HEADER, cells])
        with self.assertRaisesMessage(SpreadsheetError, 'Faltam as colunas do evento: Evento, Início'):
            read_event([('Nome', 'Matrícula'), ('Ana', '1')])
        with self.assertRaisesMessage(SpreadsheetError, 'Preencha Início'):
            read_event([HEADER, ('Semana', '', '', '', '', '', 'Ana', '', '')])


class FullSheetFlowTests(TestCase):
    def setUp(self):
        self.private = tempfile.mkdtemp()
        override = override_settings(PRIVATE_MEDIA_ROOT=Path(self.private))
        override.enable()
        self.addCleanup(override.disable)
        self.addCleanup(shutil.rmtree, self.private, True)
        self.org = make_user('org', User.Role.ORGANIZER)
        self.client.force_login(self.org)
        self.url = reverse('batch_full')

    def send(self, upload):
        return self.client.post(self.url, {'spreadsheet': upload, 'fix_case': 'on'})

    def test_emissao_completa(self):
        self.assertContains(self.client.get(reverse('batch_new')), self.url)  # opção na tela de nova emissão
        page = self.client.get(self.url)
        self.assertContains(page, 'Um evento por planilha')
        rows = [
            HEADER,
            ('Oficina de Linux', 'Curso', date(2026, 9, 1), date(2026, 9, 5), 20, date(2026, 9, 5), 'ana souza', 111,
             'ana@exemplo.com'),
            (None, None, None, None, None, None, 'Bia Lima', 222, ''),
        ]  # fmt: skip
        response = self.send(xlsx(rows))
        self.assertRedirects(response, reverse('batch_preview'))
        preview = self.client.get(reverse('batch_preview'))
        self.assertContains(preview, 'Oficina de Linux')
        self.assertContains(preview, '20 horas')
        self.assertContains(preview, 'Ana Souza')
        self.assertContains(preview, f'href="{self.url}"')  # "Enviar outra planilha" volta para esta tela
        token = self.client.session['certificate_import']['token']
        self.client.post(reverse('batch_preview'), {'token': token})
        batch = CertificateBatch.objects.get()
        self.assertEqual(
            (batch.title, batch.activity_type, batch.start_date, batch.end_date, batch.workload_hours,
             batch.issue_date),
            ('Oficina de Linux', 'COURSE', date(2026, 9, 1), date(2026, 9, 5), 20, date(2026, 9, 5)),
        )  # fmt: skip
        self.assertEqual(
            sorted(batch.certificates.values_list('participant_name', 'participant_document')),
            [('Ana Souza', '111'), ('Bia Lima', '222')],
        )

    def test_csv_e_erros_do_evento(self):
        text = 'Evento;Início;Término;Nome\nPalestra;05/09/2026;01/09/2026;Ana\n'
        response = self.send(csv_file(text))
        self.assertContains(response, 'Coluna Término: A data de término não pode ser anterior à de início.')
        response = self.send(csv_file('Evento;Início;Nome\nPalestra;01/09/2026;Ana\nOutra;01/09/2026;Bia\n'))
        self.assertContains(response, 'A linha 3 traz Evento')
        self.assertNotIn('certificate_import', self.client.session)
        response = self.send(csv_file('Evento;Início;Nome\nPalestra;01/09/2026;Ana\n'))
        self.assertRedirects(response, reverse('batch_preview'))

    def test_modelo_da_planilha_completa(self):
        response = self.client.get(reverse('batch_template') + '?modelo=completo')
        self.assertIn('modelo-planilha-completa.xlsx', response['Content-Disposition'])
        workbook = load_workbook(io.BytesIO(response.content))
        self.assertEqual(workbook.sheetnames, ['Participantes', 'Como preencher'])
        header = [c.value for c in workbook['Participantes'][1]]
        self.assertEqual(header[:7], ['Evento', 'Tipo', 'Início', 'Término', 'Carga horária', 'Data do certificado',
                                      'Nome'])  # fmt: skip
        # o modelo baixado e preenchido é aceito
        sheet = workbook['Participantes']
        sheet.append(['Palestra', '', date(2026, 9, 1), None, None, None, 'Ana', '0012', 'T1', ''])
        buffer = io.BytesIO()
        workbook.save(buffer)
        from django.core.files.uploadedfile import SimpleUploadedFile

        response = self.send(SimpleUploadedFile('planilha.xlsx', buffer.getvalue()))
        self.assertRedirects(response, reverse('batch_preview'))

    def test_permissao(self):
        self.client.force_login(make_user('ana'))
        self.assertEqual(self.client.get(self.url).status_code, 403)
