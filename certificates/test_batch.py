import csv
import io
import zipfile
from datetime import date

from django.core import mail
from django.core.files.uploadedfile import SimpleUploadedFile
from django.db import IntegrityError, transaction
from django.test import TestCase, override_settings
from django.urls import reverse
from django.utils import timezone
from openpyxl import Workbook

from accounts.models import User
from activities.test_utils import make_activity, make_user
from core.models import EmailSettings, SiteSettings
from registrations.models import Registration

from .batch_import import SpreadsheetError, column_kind, parse_participants, read_table, smart_title
from .models import Certificate, CertificateBatch
from .pdf import certificate_text

# Igual à planilha enviada pela secretaria
STUDENTS = [
    ('Nome (conforme base de acesso)', 'Matrícula', 'Turma'),
    ('Alex silva', 12546456, 'EC.7A.2026.S2'),
    ('Maria Melo', 12546457, 'EC.8A.2026.S2'),
    ('joaninha melo rego', 12546462, 'EC.10A.2026.S2'),
]


def xlsx(rows, name='alunos.xlsx'):
    workbook = Workbook()
    for row in rows:
        workbook.active.append(list(row))
    buffer = io.BytesIO()
    workbook.save(buffer)
    return SimpleUploadedFile(name, buffer.getvalue())


def csv_file(text, name='alunos.csv', encoding='utf-8'):
    return SimpleUploadedFile(name, text.encode(encoding))


class SpreadsheetParsingTests(TestCase):
    def test_planilha_no_formato_da_secretaria(self):
        rows, columns = parse_participants(read_table(xlsx(STUDENTS)))
        self.assertEqual(columns, {'name', 'document', 'group'})
        self.assertEqual([r.name for r in rows], ['Alex Silva', 'Maria Melo', 'Joaninha Melo Rego'])
        self.assertEqual(rows[0].document, '12546456')  # número do Excel vira texto, sem ".0"
        self.assertEqual(rows[0].group, 'EC.7A.2026.S2')
        self.assertEqual(rows[2].original_name, 'joaninha melo rego')
        self.assertEqual(rows[0].line, 2)
        self.assertTrue(all(r.ok for r in rows))

    def test_sem_ajuste_de_maiusculas(self):
        rows, _ = parse_participants(read_table(xlsx(STUDENTS)), fix_case=False)
        self.assertEqual(rows[0].name, 'Alex silva')

    def test_ajuste_de_maiusculas(self):
        self.assertEqual(smart_title('MARIA DA SILVA E SOUZA'), 'Maria da Silva e Souza')
        self.assertEqual(smart_title('ana-paula dos santos'), 'Ana-Paula dos Santos')
        self.assertEqual(smart_title('Ronald McDonald'), 'Ronald McDonald')
        self.assertEqual(smart_title('de souza'), 'De Souza')

    def test_reconhece_cabecalhos(self):
        self.assertEqual(column_kind('Nome completo'), 'name')
        self.assertEqual(column_kind('RA'), 'document')
        self.assertEqual(column_kind('Matricula'), 'document')
        self.assertEqual(column_kind('E-mail institucional'), 'email')
        self.assertEqual(column_kind('Nome da turma'), 'group')
        self.assertIsNone(column_kind('Observação'))

    def test_problemas_por_linha(self):
        table = [
            ['Título qualquer'],  # o cabeçalho pode não estar na primeira linha
            ['RA', 'Nome', 'Email'],
            ['1', 'Ana', 'ana@exemplo.com'],
            ['1', 'Ana Outra', ''],  # matrícula repetida
            ['', '', ''],  # linha vazia: ignorada
            ['2', '', 'x@exemplo.com'],  # sem nome
            ['3', 'Bia', 'bia@'],  # e-mail inválido
            ['', 'Caio', ''],  # sem matrícula: aviso, mas emite
        ]
        rows, columns = parse_participants(table)
        self.assertEqual(columns, {'name', 'document', 'email'})
        problems = {r.line: (r.errors, r.warnings) for r in rows}
        self.assertEqual(problems[3], ([], []))
        self.assertEqual(problems[4][0], ['Repetido (linha 3)'])
        self.assertNotIn(5, problems)
        self.assertEqual(problems[6][0], ['Sem nome'])
        self.assertEqual(problems[7][0], ['E-mail inválido'])
        self.assertEqual(problems[8], ([], ['Sem matrícula']))

    def test_externos_sem_matricula_nao_sao_repetidos(self):
        table = [
            ['Nome', 'Matrícula', 'Turma'],
            ['Ana Lima', 'Externo', 'Externo'],
            ['LEANDRO VIRGINIO DOS SANTOS', 'Externo', 'Externo'],
            ['Bia Souza', '-', 'Externo'],
            ['Ana Lima', 'externo', 'Externo'],  # mesma pessoa duas vezes: continua sendo repetido
            ['Caio', 'A12345', 'EC.1A'],  # matrícula com letras e números é aceita
        ]
        rows, _ = parse_participants(table)
        by_line = {r.line: r for r in rows}
        self.assertEqual(by_line[3].name, 'Leandro Virginio dos Santos')
        self.assertEqual((by_line[3].document, by_line[3].errors), ('', []))
        self.assertEqual(by_line[3].warnings, ['Sem matrícula ("Externo")'])
        self.assertEqual(by_line[4].errors, [])
        self.assertEqual(by_line[5].errors, ['Repetido (linha 2)'])
        self.assertEqual((by_line[6].document, by_line[6].warnings), ('A12345', []))

    def test_csv_com_ponto_e_virgula_e_acentos_do_excel(self):
        upload = csv_file('Nome;Matrícula\nJoão Conceição;123\n', encoding='cp1252')
        rows, _ = parse_participants(read_table(upload))
        self.assertEqual((rows[0].name, rows[0].document), ('João Conceição', '123'))
        rows, _ = parse_participants(read_table(csv_file('Nome,RA\nAna,9\n')))
        self.assertEqual(rows[0].document, '9')

    def test_arquivos_rejeitados(self):
        with self.assertRaisesMessage(SpreadsheetError, '.xls (Excel antigo)'):
            read_table(SimpleUploadedFile('alunos.xls', b'x'))
        with self.assertRaisesMessage(SpreadsheetError, 'xlsx válido'):
            read_table(SimpleUploadedFile('alunos.xlsx', b'nao e excel'))
        with self.assertRaisesMessage(SpreadsheetError, 'coluna com o nome'):
            parse_participants([['Matrícula', 'Turma'], ['1', 'A']])
        with self.assertRaisesMessage(SpreadsheetError, 'não tem participantes'):
            parse_participants([['Nome'], [None]])


class BatchFlowTests(TestCase):
    def setUp(self):
        self.org = make_user('org', User.Role.ORGANIZER, first_name='Manoel', last_name='Prof')
        self.client.force_login(self.org)

    def event_data(self, **changes):
        data = {
            'title': 'Semana de Engenharia',
            'activity_type': 'EVENT',
            'start_date': '2026-08-10',
            'end_date': '2026-08-14',
            'workload_hours': '12',
            'issue_date': '2026-08-14',
            'fix_case': 'on',
        }
        data.update(changes)
        return data

    def upload(self, rows=STUDENTS, **changes):
        return self.client.post(reverse('batch_new'), {**self.event_data(**changes), 'spreadsheet': xlsx(rows)})

    def confirm(self):
        token = self.client.session['certificate_import']['token']
        return self.client.post(reverse('batch_preview'), {'token': token}, follow=True)

    def issue(self, rows=STUDENTS):
        self.upload(rows)
        self.confirm()
        return CertificateBatch.objects.latest('created_at')

    def test_fluxo_completo(self):
        response = self.upload()
        self.assertRedirects(response, reverse('batch_preview'))
        preview = self.client.get(reverse('batch_preview'))
        self.assertContains(preview, 'Joaninha Melo Rego')
        self.assertContains(preview, 'na planilha: joaninha melo rego')
        self.assertContains(preview, 'Emitir 3 certificados')
        self.assertFalse(Certificate.objects.exists())  # a prévia não grava nada

        response = self.confirm()
        batch = CertificateBatch.objects.get()
        self.assertRedirects(response, reverse('batch_detail', args=[batch.pk]))
        self.assertContains(response, '3 certificados emitidos.')
        self.assertEqual(batch.created_by, self.org)
        self.assertEqual(
            (batch.start_date, batch.end_date, batch.workload_hours), (date(2026, 8, 10), date(2026, 8, 14), 12)
        )

        alex = Certificate.objects.get(participant_document='12546456')
        self.assertEqual(alex.participant_name, 'Alex Silva')
        self.assertEqual(alex.participant_group, 'EC.7A.2026.S2')
        self.assertEqual(alex.activity_title, 'Semana de Engenharia')
        self.assertEqual(alex.activity_type, 'Evento')
        self.assertIsNone(alex.registration)
        self.assertIsNone(alex.attendance_percent)
        self.assertEqual(timezone.localtime(alex.start_date).date(), date(2026, 8, 10))
        self.assertEqual(timezone.localtime(alex.end_date).date(), date(2026, 8, 14))

        # a prévia não pode ser usada duas vezes
        self.assertNotIn('certificate_import', self.client.session)
        self.assertEqual(Certificate.objects.count(), 3)

    def test_data_retroativa(self):
        self.issue()
        alex = Certificate.objects.get(participant_document='12546456')
        self.assertEqual(alex.document_date, date(2026, 8, 14))
        self.assertEqual(alex.certificate_date, date(2026, 8, 14))
        self.assertEqual(alex.issued_at.date(), timezone.now().date())  # quando foi gerado de fato
        page = self.client.get(reverse('certificate_verify_code', args=[alex.code]))
        self.assertContains(page, '14/08/2026</dd>')

    def test_data_do_certificado_valida(self):
        tomorrow = (timezone.localdate() + timezone.timedelta(days=1)).isoformat()
        self.assertContains(self.upload(issue_date=tomorrow), 'não pode ser no futuro')
        self.assertContains(self.upload(issue_date='2026-08-01'), 'anterior ao início do evento')
        form = self.client.get(reverse('batch_new'))
        self.assertContains(form, f'value="{timezone.localdate().isoformat()}"')  # padrão: hoje
        self.assertContains(form, 'Usar o último dia do evento')

    def test_validacao_publica_sem_frequencia(self):
        self.issue()
        alex = Certificate.objects.get(participant_document='12546456')
        self.client.logout()
        page = self.client.get(reverse('certificate_verify_code', args=[alex.code]))
        self.assertContains(page, 'Alex Silva')
        self.assertContains(page, 'Semana de Engenharia')
        self.assertNotContains(page, 'Frequência')

    def test_pdf_do_certificado_avulso(self):
        self.issue()
        alex = Certificate.objects.get(participant_document='12546456')
        response = self.client.get(reverse('certificate_download', args=[alex.code]))
        self.assertEqual(response['Content-Type'], 'application/pdf')
        self.assertTrue(response.content.startswith(b'%PDF'))
        self.assertEqual(certificate_text(alex, SiteSettings.load())['date_line'], 'São Paulo, 14 de agosto de 2026.')
        # outro organizador não baixa; administrador baixa
        self.client.force_login(make_user('outro', User.Role.ORGANIZER))
        self.assertEqual(self.client.get(reverse('certificate_download', args=[alex.code])).status_code, 404)
        self.client.force_login(make_user('adm', User.Role.ADMIN))
        self.assertEqual(self.client.get(reverse('certificate_download', args=[alex.code])).status_code, 200)

    def test_linhas_com_problema_ficam_de_fora(self):
        rows = [('Nome', 'Matrícula'), ('Ana', '1'), ('Ana de novo', '1'), ('', '2')]
        self.upload(rows)
        preview = self.client.get(reverse('batch_preview'))
        self.assertContains(preview, 'Emitir 1 certificado<')
        self.assertContains(preview, 'Repetido (linha 2)')
        response = self.confirm()
        self.assertContains(response, '1 certificado emitido. 2 linhas com problema ficaram de fora.')
        self.assertEqual(Certificate.objects.count(), 1)

    def test_planilha_ruim_volta_ao_formulario_com_os_dados(self):
        response = self.client.post(
            reverse('batch_new'), {**self.event_data(), 'spreadsheet': SimpleUploadedFile('a.xls', b'x')}
        )
        self.assertEqual(response.status_code, 200)
        self.assertContains(response, 'Excel antigo')
        self.assertContains(response, 'value="Semana de Engenharia"')

    def test_validacoes_do_evento(self):
        response = self.upload(end_date='2026-08-01')
        self.assertContains(response, 'não pode ser anterior')
        response = self.upload(workload_hours='0')
        self.assertContains(response, 'entre 1 e 2000 horas')
        self.assertNotIn('certificate_import', self.client.session)

    def test_adicionar_participantes_ignora_quem_ja_tem(self):
        batch = self.issue()
        more = [('Nome', 'RA'), ('ALEX SILVA', '12546456'), ('Novo Aluno', '999')]
        self.client.post(reverse('batch_add', args=[batch.pk]), {'spreadsheet': xlsx(more), 'fix_case': 'on'})
        preview = self.client.get(reverse('batch_preview'))
        self.assertContains(preview, 'Já tem certificado neste lote')
        self.assertContains(preview, 'Emitir 1 certificado<')
        self.confirm()
        self.assertEqual(batch.certificates.count(), 4)
        self.assertTrue(batch.certificates.filter(participant_name='Novo Aluno', activity_title=batch.title).exists())

    def test_zip_e_csv(self):
        batch = self.issue()
        response = self.client.get(reverse('batch_zip', args=[batch.pk]))
        archive = zipfile.ZipFile(io.BytesIO(b''.join(response.streaming_content)))
        names = sorted(archive.namelist())
        self.assertEqual(
            names, ['12546456-alex-silva.pdf', '12546457-maria-melo.pdf', '12546462-joaninha-melo-rego.pdf']
        )
        self.assertTrue(archive.read(names[0]).startswith(b'%PDF'))

        response = self.client.get(reverse('batch_csv', args=[batch.pk]))
        text = response.content.decode('utf-8-sig')
        lines = list(csv.reader(io.StringIO(text), delimiter=';'))
        self.assertEqual(lines[0][:5], ['Nome', 'Matrícula', 'Turma', 'E-mail', 'Código'])
        alex = batch.certificates.get(participant_document='12546456')
        self.assertIn(['Alex Silva', '12546456', 'EC.7A.2026.S2', '', alex.code], [line[:5] for line in lines])
        self.assertTrue(lines[1][5].endswith(f'/certificados/validar/{lines[1][4]}/'))

    def test_editar_atualiza_os_certificados(self):
        batch = self.issue()
        response = self.client.post(
            reverse('batch_edit', args=[batch.pk]),
            {'title': 'Semana da Engenharia 2026', 'activity_type': 'EVENT', 'start_date': '2026-08-11',
             'end_date': '2026-08-11', 'workload_hours': '', 'issue_date': '2026-08-20'},
        )  # fmt: skip
        self.assertRedirects(response, reverse('batch_detail', args=[batch.pk]))
        certificate = batch.certificates.first()
        self.assertEqual(certificate.activity_title, 'Semana da Engenharia 2026')
        self.assertIsNone(certificate.workload_hours)
        self.assertEqual(timezone.localtime(certificate.end_date).date(), date(2026, 8, 11))
        self.assertEqual(certificate.certificate_date, date(2026, 8, 20))

    def test_revogar_certificado_e_reemitir(self):
        batch = self.issue()
        alex = batch.certificates.get(participant_document='12546456')
        url = reverse('batch_certificate_revoke', args=[batch.pk, alex.code])
        self.assertContains(self.client.get(url), 'Revogar certificado')
        self.assertContains(self.client.post(url, {'reason': 'x'}), 'pelo menos 5 caracteres')  # motivo obrigatório
        self.client.post(url, {'reason': 'Nome digitado errado'})
        alex.refresh_from_db()
        self.assertTrue(alex.is_revoked)
        self.assertEqual((alex.revoked_by, alex.revoked_reason), (self.org, 'Nome digitado errado'))
        self.assertEqual(batch.certificates.count(), 3)  # nada é apagado

        page = self.client.get(reverse('certificate_verify_code', args=[alex.code]))
        self.assertContains(page, 'Certificado revogado')
        self.assertContains(page, 'Nome digitado errado')
        self.assertNotContains(page, 'Certificado autêntico')
        response = self.client.get(reverse('certificate_download', args=[alex.code]), follow=True)
        self.assertContains(response, 'foi revogado e não pode mais ser baixado')

        # a mesma matrícula pode ser emitida de novo, com código novo
        self.client.post(
            reverse('batch_add', args=[batch.pk]),
            {'spreadsheet': xlsx([('Nome', 'RA'), ('Alex da Silva', '12546456')]), 'fix_case': 'on'},
        )
        self.confirm()
        new = batch.certificates.get(participant_document='12546456', revoked_at__isnull=True)
        self.assertNotEqual(new.code, alex.code)
        zip_names = zipfile.ZipFile(
            io.BytesIO(b''.join(self.client.get(reverse('batch_zip', args=[batch.pk])).streaming_content))
        ).namelist()
        self.assertEqual(len(zip_names), 3)  # o revogado fica fora do ZIP
        detail = self.client.get(reverse('batch_detail', args=[batch.pk]))
        self.assertContains(detail, 'Revogados (1)')

    def test_cancelar_lote(self):
        batch = self.issue()
        url = reverse('batch_cancel', args=[batch.pk])
        self.assertContains(self.client.get(url), '3 certificados deste lote serão revogados')
        self.client.post(url, {'reason': 'Evento lançado em duplicidade'})
        batch.refresh_from_db()
        self.assertTrue(batch.is_cancelled)
        self.assertEqual(batch.certificates.filter(revoked_at__isnull=False).count(), 3)
        self.assertEqual(Certificate.objects.count(), 3)
        detail = self.client.get(reverse('batch_detail', args=[batch.pk]))
        self.assertContains(detail, 'Emissão cancelada em')
        self.assertNotContains(detail, reverse('batch_add', args=[batch.pk]))
        for name in ('batch_add', 'batch_edit', 'batch_cancel'):
            response = self.client.get(reverse(name, args=[batch.pk]), follow=True)
            self.assertContains(response, 'foi cancelada e não pode mais ser alterada')
        self.assertContains(self.client.get(reverse('batch_list')), 'Cancelada')

    def test_permissoes(self):
        batch = self.issue()
        urls = [reverse('batch_detail', args=[batch.pk]), reverse('batch_zip', args=[batch.pk]),
                reverse('batch_csv', args=[batch.pk]), reverse('batch_add', args=[batch.pk]),
                reverse('batch_edit', args=[batch.pk]), reverse('batch_cancel', args=[batch.pk])]  # fmt: skip
        self.client.force_login(make_user('ana'))
        for url in [reverse('batch_list'), reverse('batch_new'), reverse('batch_template'), *urls]:
            self.assertEqual(self.client.get(url).status_code, 403, url)
        self.client.force_login(make_user('outro', User.Role.ORGANIZER))
        for url in urls:
            self.assertEqual(self.client.get(url).status_code, 403, url)
        self.assertNotContains(self.client.get(reverse('batch_list')), 'Semana de Engenharia')
        admin = make_user('adm', User.Role.ADMIN)
        self.client.force_login(admin)
        self.assertContains(self.client.get(reverse('batch_list')), 'Semana de Engenharia')
        self.assertEqual(self.client.get(urls[0]).status_code, 200)

    def test_planilha_modelo(self):
        from openpyxl import load_workbook

        response = self.client.get(reverse('batch_template'))
        sheet = load_workbook(io.BytesIO(response.content)).active
        self.assertEqual([c.value for c in sheet[1]], ['Nome', 'Matrícula', 'Turma', 'E-mail'])

    def test_painel_tem_atalho(self):
        self.assertContains(self.client.get(reverse('organizer_dashboard')), reverse('batch_list'))


@override_settings(EMAIL_BACKEND='django.core.mail.backends.locmem.EmailBackend')
class BatchEmailTests(TestCase):
    def setUp(self):
        self.org = make_user('org', User.Role.ORGANIZER)
        self.client.force_login(self.org)
        rows = [
            ('Nome', 'RA', 'E-mail'),
            ('Ana', '1', 'Ana@Exemplo.com'),
            ('Bia', '2', ''),
            ('Caio', '3', 'caio@exemplo.com'),
        ]
        self.client.post(
            reverse('batch_new'),
            {'title': 'Palestra', 'activity_type': 'EVENT', 'start_date': '2026-09-01', 'end_date': '2026-09-01',
             'issue_date': '2026-09-01',
             'spreadsheet': xlsx(rows)},
        )  # fmt: skip
        token = self.client.session['certificate_import']['token']
        self.client.post(reverse('batch_preview'), {'token': token})
        self.batch = CertificateBatch.objects.get()
        self.url = reverse('batch_email', args=[self.batch.pk])

    def enable(self):
        config = EmailSettings.load()
        config.enabled = True
        config.save()

    def test_envio_desligado(self):
        response = self.client.post(self.url, follow=True)
        self.assertContains(response, 'O envio de e-mails está desligado')
        self.assertEqual(len(mail.outbox), 0)

    def test_envia_com_pdf_so_para_pendentes(self):
        self.enable()
        detail = self.client.get(reverse('batch_detail', args=[self.batch.pk]))
        self.assertContains(detail, 'Enviar para 2 pendentes')
        response = self.client.post(self.url, follow=True)
        self.assertContains(response, '2 certificados enviados por e-mail.')
        self.assertEqual(sorted(m.to[0] for m in mail.outbox), ['ana@exemplo.com', 'caio@exemplo.com'])
        message = next(m for m in mail.outbox if m.to == ['ana@exemplo.com'])
        self.assertIn('Palestra', message.subject)
        self.assertIn('Olá, Ana.', message.body)
        filename, content, mimetype = message.attachments[0]
        self.assertEqual((filename, mimetype), ('1-ana.pdf', 'application/pdf'))
        self.assertTrue(content.startswith(b'%PDF'))
        self.assertEqual(self.batch.certificates.filter(emailed_at__isnull=False).count(), 2)

        mail.outbox.clear()
        response = self.client.post(self.url, follow=True)
        self.assertContains(response, 'Não havia certificados pendentes')
        self.client.post(self.url, {'resend': '1'})
        self.assertEqual(len(mail.outbox), 2)


class CertificateConstraintTests(TestCase):
    def test_certificado_e_de_inscricao_ou_de_lote(self):
        org = make_user('org', User.Role.ORGANIZER)
        registration = Registration.objects.create(user=make_user('ana'), activity=make_activity(org))
        batch = CertificateBatch.objects.create(
            title='X', start_date=date(2026, 1, 1), end_date=date(2026, 1, 1), created_by=org
        )
        base = {'participant_name': 'Ana', 'activity_title': 'X', 'activity_type': 'Evento',
                'start_date': timezone.now(), 'end_date': timezone.now()}  # fmt: skip
        for kwargs in ({}, {'registration': registration, 'batch': batch}):
            with self.assertRaises(IntegrityError), transaction.atomic():
                Certificate.objects.create(**base, **kwargs)
        Certificate.objects.create(**base, batch=batch, participant_document='1')
        with self.assertRaises(IntegrityError), transaction.atomic():
            Certificate.objects.create(**base, batch=batch, participant_document='1')
