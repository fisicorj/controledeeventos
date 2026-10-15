"""Telas da emissão avulsa de certificados (evento que não foi cadastrado no sistema)."""

import io
import json
import re
import secrets
import time
from datetime import date, timedelta
from pathlib import Path

from django.conf import settings
from django.contrib import messages
from django.contrib.auth.decorators import login_required
from django.core.exceptions import PermissionDenied
from django.db import transaction
from django.db.models import Count, Q
from django.http import FileResponse, Http404, HttpResponse
from django.shortcuts import get_object_or_404, redirect, render
from django.urls import reverse
from django.utils import timezone
from django.utils.text import slugify
from django.views.decorators.http import require_POST

from accounts.reauth import require_recent_auth
from core import outbox
from core.csvsafe import excel_writer
from core.models import EmailSettings, OutgoingEmail
from core.pagination import paginate
from core.permissions import can_manage_batch

from . import zips
from .batch import (
    active,
    batches_for,
    build_zip,
    cancel_batch,
    email_group,
    existing_keys,
    issue_batch,
    queue_emails,
    revoke,
    update_batch,
)
from .batch_import import EVENT_COLUMNS, SpreadsheetError, parse_participants, read_event, read_table
from .forms import BatchForm, FullSpreadsheetForm, ReasonForm, SpreadsheetForm
from .models import CertificateBatch, ZipExport

SESSION_KEY = 'certificate_import'
FORBIDDEN = 'Acesso restrito a organizadores e administradores.'


def _get_batch(request, pk):
    batch = get_object_or_404(CertificateBatch.objects.select_related('created_by'), pk=pk)
    return batch if can_manage_batch(request.user, batch) else None


def _closed(request, batch):
    """Lote cancelado não aceita mais alterações. Devolve o redirecionamento, se for o caso."""
    if batch.is_cancelled:
        messages.error(request, 'Esta emissão foi cancelada e não pode mais ser alterada.')
        return redirect('batch_detail', pk=batch.pk)
    return None


def _filename(batch, suffix):
    return f'{slugify(batch.title)[:50] or "certificados"}-{suffix}'


@login_required
def batch_list(request):
    if not can_manage_batch(request.user):
        raise PermissionDenied(FORBIDDEN)
    batches = batches_for(request.user)
    query = request.GET.get('q', '').strip()
    if query:
        batches = batches.filter(title__icontains=query)
    batches = batches.annotate(
        total=Count('certificates', filter=Q(certificates__revoked_at__isnull=True)),
        revoked=Count('certificates', filter=Q(certificates__revoked_at__isnull=False)),
    )
    page, querystring = paginate(request, batches, 20)
    context = {'batches': page.object_list, 'page': page, 'querystring': querystring, 'query': query}
    return render(request, 'certificates/batch_list.html', context)


def _read_spreadsheet(form, batch=None):
    """Lê e valida a planilha. Devolve as linhas ou None (com o erro anotado no formulário)."""
    try:
        table = read_table(form.cleaned_data['spreadsheet'])
        rows, columns = parse_participants(table, fix_case=form.cleaned_data['fix_case'], existing=existing_keys(batch))
    except SpreadsheetError as exc:
        form.add_error('spreadsheet', str(exc))
        return None, None
    return rows, columns


# As linhas da prévia (até 3.000) ficam num arquivo temporário, não na sessão: a sessão guarda só o essencial.
PREVIEW_MAX_AGE = 24 * 3600


def _preview_dir():
    return Path(settings.PRIVATE_MEDIA_ROOT) / 'tmp' / 'importacoes'


def _preview_file(token):
    return _preview_dir() / f'{token}.json'


def _clean_old_previews():
    limit = time.time() - PREVIEW_MAX_AGE
    for path in _preview_dir().glob('*.json'):
        try:
            if path.stat().st_mtime < limit:
                path.unlink()
        except OSError:
            pass


def _store_preview(request, rows, columns, filename, batch=None, event=None, mode='form'):
    _drop_preview(request)
    token = secrets.token_urlsafe(16)
    _preview_dir().mkdir(parents=True, exist_ok=True)
    _clean_old_previews()
    _preview_file(token).write_text(json.dumps([row.to_dict() for row in rows], ensure_ascii=False), 'utf-8')
    request.session[SESSION_KEY] = {
        'token': token,
        'batch_id': batch.pk if batch else None,
        'event': event,
        'filename': filename,
        'columns': sorted(columns),
        'mode': mode,  # "form": dados do evento digitados; "full": vieram da planilha completa
    }
    return redirect('batch_preview')


def _load_preview(request):
    """Dados da prévia com as linhas lidas do arquivo; None se não houver ou se o arquivo sumiu."""
    data = request.session.get(SESSION_KEY)
    if not data or not re.fullmatch(r'[\w-]+', data.get('token', '')):
        return None
    try:
        rows = json.loads(_preview_file(data['token']).read_text('utf-8'))
    except (OSError, ValueError):
        request.session.pop(SESSION_KEY, None)
        return None
    return {**data, 'rows': rows}


def _drop_preview(request):
    data = request.session.pop(SESSION_KEY, None)
    if data and re.fullmatch(r'[\w-]+', data.get('token', '')):
        _preview_file(data['token']).unlink(missing_ok=True)


@login_required
def batch_new(request):
    if not can_manage_batch(request.user):
        raise PermissionDenied(FORBIDDEN)
    # Voltando da prévia ("Enviar outra planilha"), o formulário vem preenchido com o evento digitado.
    previous = request.session.get(SESSION_KEY) or {}
    initial = previous.get('event') if not previous.get('batch_id') and previous.get('mode') != 'full' else None
    event_form = BatchForm(request.POST or None, initial=initial)
    sheet_form = SpreadsheetForm(request.POST or None, request.FILES or None)
    if request.method == 'POST' and event_form.is_valid() and sheet_form.is_valid():
        rows, columns = _read_spreadsheet(sheet_form)
        if rows is not None:
            event = _event_session(event_form.cleaned_data)
            return _store_preview(request, rows, columns, sheet_form.cleaned_data['spreadsheet'].name, event=event)
    return render(
        request, 'certificates/batch_form.html', {'event_form': event_form, 'sheet_form': sheet_form, 'batch': None}
    )


def _event_session(cleaned):
    return {key: value.isoformat() if isinstance(value, date) else value for key, value in cleaned.items()}


@login_required
def batch_full(request):
    """Nova emissão com a planilha completa: evento e participantes no mesmo arquivo (um evento por planilha)."""
    if not can_manage_batch(request.user):
        raise PermissionDenied(FORBIDDEN)
    sheet_form = FullSpreadsheetForm(request.POST or None, request.FILES or None)
    if request.method == 'POST' and sheet_form.is_valid():
        upload = sheet_form.cleaned_data['spreadsheet']
        try:
            table = read_table(upload)
            event = read_event(table)
            # Os dados lidos passam pelas mesmas validações do formulário digitado (datas, carga horária...).
            event_form = BatchForm(
                data={key: '' if value is None else value for key, value in _event_session(event).items()}
            )
            if not event_form.is_valid():
                labels = {**EVENT_COLUMNS, 'start_date': 'Início', 'end_date': 'Término'}
                for field, errors in event_form.errors.items():
                    for error in errors:
                        prefix = f'Coluna {labels[field]}: ' if field in labels else ''
                        sheet_form.add_error('spreadsheet', prefix + error)
                raise SpreadsheetError('')
            rows, columns = parse_participants(table, fix_case=sheet_form.cleaned_data['fix_case'])
        except SpreadsheetError as exc:
            if str(exc):
                sheet_form.add_error('spreadsheet', str(exc))
        else:
            return _store_preview(
                request, rows, columns, upload.name, event=_event_session(event_form.cleaned_data), mode='full'
            )
    return render(
        request,
        'certificates/batch_full.html',
        {'sheet_form': sheet_form, 'columns': EVENT_COLUMNS},
    )


@login_required
def batch_add(request, pk):
    batch = _get_batch(request, pk)
    if batch is None:
        raise PermissionDenied(FORBIDDEN)
    if closed := _closed(request, batch):
        return closed
    sheet_form = SpreadsheetForm(request.POST or None, request.FILES or None)
    if request.method == 'POST' and sheet_form.is_valid():
        rows, columns = _read_spreadsheet(sheet_form, batch)
        if rows is not None:
            return _store_preview(request, rows, columns, sheet_form.cleaned_data['spreadsheet'].name, batch=batch)
    return render(request, 'certificates/batch_form.html', {'sheet_form': sheet_form, 'batch': batch})


@login_required
def batch_preview(request):
    if not can_manage_batch(request.user):
        raise PermissionDenied(FORBIDDEN)
    data = _load_preview(request)
    if not data:
        messages.info(request, 'Envie a planilha para ver a prévia.')
        return redirect('batch_list')
    batch = None
    if data['batch_id']:
        batch = _get_batch(request, data['batch_id'])
        if batch is None:
            raise PermissionDenied(FORBIDDEN)
        if closed := _closed(request, batch):
            _drop_preview(request)
            return closed

    rows = data['rows']
    valid = [r for r in rows if not r['errors']]

    if request.method == 'POST':
        return _issue_preview(request, data, batch, rows, valid)
    return render(request, 'certificates/batch_preview.html', _preview_context(data, batch, rows, valid))


def _issue_preview(request, data, batch, rows, valid):
    """Confirmação da prévia: cria o lote (se for novo) e emite as linhas válidas."""
    if request.POST.get('token') != data['token']:
        messages.error(request, 'Esta prévia já foi usada ou expirou. Envie a planilha de novo.')
        return redirect('batch_list')
    if not valid:
        messages.error(request, 'Nenhuma linha válida para emitir.')
        return redirect('batch_preview')
    with transaction.atomic():
        if batch is None:
            event = data['event']
            batch = CertificateBatch.objects.create(
                title=event['title'],
                activity_type=event['activity_type'],
                start_date=date.fromisoformat(event['start_date']),
                end_date=date.fromisoformat(event['end_date']),
                workload_hours=event['workload_hours'],
                issue_date=date.fromisoformat(event['issue_date']),
                created_by=request.user,
            )
        created = issue_batch(batch, rows, request=request)
    _drop_preview(request)
    skipped = len(rows) - len(created)
    text = f'{len(created)} certificado{"s" if len(created) != 1 else ""} emitido{"s" if len(created) != 1 else ""}.'
    if skipped:
        plural = skipped != 1
        text += f' {skipped} linha{"s" if plural else ""} com problema {"ficaram" if plural else "ficou"} de fora.'
    messages.success(request, text)
    return redirect('batch_detail', pk=batch.pk)


def _preview_context(data, batch, rows, valid):
    context = {
        'data': data,
        'batch': batch,
        'rows': rows,
        'valid_count': len(valid),
        'error_count': len(rows) - len(valid),
        'warning_count': sum(1 for r in valid if r['warnings']),
        'renamed_count': sum(1 for r in rows if r['name'] != r['original_name']),
        'has_email': 'email' in data['columns'],
        'has_group': 'group' in data['columns'],
        'has_document': 'document' in data['columns'],
    }
    if batch is None:
        event = data['event']
        context['event'] = {
            **event,
            'start': date.fromisoformat(event['start_date']),
            'end': date.fromisoformat(event['end_date']),
            'issue': date.fromisoformat(event['issue_date']),
            'type_label': dict(CertificateBatch._meta.get_field('activity_type').choices)[event['activity_type']],
        }
    return context


@login_required
def batch_detail(request, pk):
    batch = _get_batch(request, pk)
    if batch is None:
        raise PermissionDenied(FORBIDDEN)
    everything = list(batch.certificates.select_related('revoked_by').order_by('participant_name'))
    certificates = [c for c in everything if not c.is_revoked]
    # Lotes grandes (até 3.000 pessoas): a tabela mostra 100 por página, com busca por nome, matrícula ou turma.
    query = request.GET.get('busca', '').strip().lower()
    shown = certificates
    if query:
        shown = [
            c
            for c in certificates
            if query in c.participant_name.lower()
            or query in c.participant_document.lower()
            or query in c.participant_group.lower()
        ]
    page, querystring = paginate(request, shown, 100)
    revoked = sorted((c for c in everything if c.is_revoked), key=lambda c: c.revoked_at, reverse=True)
    with_email = [c for c in certificates if c.participant_email]
    queue_status = outbox.group_status(email_group(batch))[email_group(batch)]
    last_failure = (
        OutgoingEmail.objects.filter(group=email_group(batch), status=OutgoingEmail.Status.FAILED)
        .order_by('-id')
        .values_list('last_error', flat=True)
        .first()
    )
    return render(
        request,
        'certificates/batch_detail.html',
        {
            'batch': batch,
            'certificates': certificates,
            'page': page,
            'querystring': querystring,
            'query': request.GET.get('busca', '').strip(),
            'revoked': revoked,
            'events': batch.events.select_related('actor')[:15],
            'with_email': len(with_email),
            'pending_email': sum(1 for c in with_email if not c.emailed_at),
            'email_enabled': EmailSettings.load().enabled,
            'queue': queue_status,
            'last_failure': last_failure,
            'has_group': any(c.participant_group for c in everything),
            'zip_direct': len(certificates) <= zips.ZIP_DIRECT_LIMIT,
            'zip_export': zips.latest(batch),
        },
    )


@login_required
@require_recent_auth
def batch_edit(request, pk):
    batch = _get_batch(request, pk)
    if batch is None:
        raise PermissionDenied(FORBIDDEN)
    if closed := _closed(request, batch):
        return closed
    form = BatchForm(request.POST or None, instance=batch)
    if form.is_valid():
        with transaction.atomic():
            batch = form.save()
            changed = update_batch(batch, request=request)
        messages.success(
            request,
            f'Dados do evento salvos. {changed} certificado{"s" if changed != 1 else ""} '
            f'atualizado{"s" if changed != 1 else ""} (as mudanças ficam registradas na auditoria).',
        )
        return redirect('batch_detail', pk=batch.pk)
    return render(request, 'certificates/batch_edit.html', {'form': form, 'batch': batch})


@login_required
@require_recent_auth
def batch_cancel(request, pk):
    """Cancela a emissão inteira: revoga todos os certificados. Nada é apagado."""
    batch = _get_batch(request, pk)
    if batch is None:
        raise PermissionDenied(FORBIDDEN)
    if closed := _closed(request, batch):
        return closed
    total = active(batch).count()
    form = ReasonForm(request.POST if request.method == 'POST' else None)
    if form.is_valid():
        revoked = cancel_batch(batch, form.cleaned_data['reason'], request)
        messages.success(
            request,
            f'Emissão "{batch.title}" cancelada. {len(revoked)} certificado{"s" if len(revoked) != 1 else ""} '
            f'revogado{"s" if len(revoked) != 1 else ""}.',
        )
        return redirect('batch_detail', pk=batch.pk)
    return render(
        request,
        'certificates/revoke.html',
        {
            'form': form,
            'batch': batch,
            'heading': 'Cancelar emissão avulsa',
            'subject': batch.title,
            'consequence': f'{total} certificado{"s" if total != 1 else ""} deste lote '
            f'{"serão revogados" if total != 1 else "será revogado"}.',
            'button': 'Cancelar emissão e revogar certificados',
            'back_url': reverse('batch_detail', args=[batch.pk]),
        },
    )


@login_required
@require_recent_auth
def batch_certificate_revoke(request, pk, code):
    batch = _get_batch(request, pk)
    if batch is None:
        raise PermissionDenied(FORBIDDEN)
    certificate = get_object_or_404(batch.certificates, code=code)
    if certificate.is_revoked:
        messages.info(request, 'Este certificado já estava revogado.')
        return redirect('batch_detail', pk=batch.pk)
    form = ReasonForm(request.POST if request.method == 'POST' else None)
    if form.is_valid():
        revoke([certificate], form.cleaned_data['reason'], request)
        messages.success(request, f'Certificado de {certificate.participant_name} ({certificate.code}) revogado.')
        return redirect('batch_detail', pk=batch.pk)
    return render(
        request,
        'certificates/revoke.html',
        {
            'form': form,
            'batch': batch,
            'certificate': certificate,
            'heading': 'Revogar certificado',
            'subject': f'{certificate.participant_name} · {certificate.code}',
            'consequence': 'O certificado continua guardado, mas não pode mais ser baixado nem enviado.',
            'button': 'Revogar certificado',
            'back_url': reverse('batch_detail', args=[batch.pk]),
        },
    )


@login_required
def batch_zip(request, pk):
    batch = _get_batch(request, pk)
    if batch is None:
        raise PermissionDenied(FORBIDDEN)
    if active(batch).count() > zips.ZIP_DIRECT_LIMIT:  # lote grande: monta em segundo plano (DES-02)
        messages.info(request, 'Este lote é grande: use "Gerar ZIP" e baixe quando ficar pronto.')
        return redirect('batch_detail', pk=batch.pk)
    return FileResponse(
        build_zip(batch, request=request), as_attachment=True, filename=_filename(batch, 'certificados.zip')
    )


@login_required
@require_POST
def batch_zip_request(request, pk):
    batch = _get_batch(request, pk)
    if batch is None:
        raise PermissionDenied(FORBIDDEN)
    zips.request_export(batch, request.user)
    messages.success(request, 'O ZIP está sendo gerado. Esta página atualiza sozinha; o link aparece aqui.')
    return redirect('batch_detail', pk=batch.pk)


@login_required
def batch_zip_file(request, pk, export_pk):
    batch = _get_batch(request, pk)
    if batch is None:
        raise PermissionDenied(FORBIDDEN)
    export = get_object_or_404(batch.zip_exports, pk=export_pk, status=ZipExport.Status.READY)
    if not export.file or export.created_at < timezone.now() - timedelta(hours=zips.ZIP_KEEP_HOURS):
        raise Http404
    return FileResponse(export.file.open('rb'), as_attachment=True, filename=_filename(batch, 'certificados.zip'))


@login_required
def batch_csv(request, pk):
    batch = _get_batch(request, pk)
    if batch is None:
        raise PermissionDenied(FORBIDDEN)
    response = HttpResponse(content_type='text/csv; charset=utf-8')
    response['Content-Disposition'] = f'attachment; filename="{_filename(batch, "certificados")}.csv"'
    response.write('﻿')  # BOM: o Excel abre com acentos corretos
    writer = excel_writer(response)
    writer.writerow(
        [
            'Nome',
            'Matrícula',
            'Turma',
            'E-mail',
            'Código',
            'Link de validação',
            'Data do certificado',
            'Gerado no sistema em',
            'Enviado por e-mail em',
            'Situação',
            'Revogado em',
            'Motivo da revogação',
        ]
    )
    for c in batch.certificates.order_by('participant_name'):
        writer.writerow([
            c.participant_name,
            c.participant_document,
            c.participant_group,
            c.participant_email,
            c.code,
            request.build_absolute_uri(reverse('certificate_verify_code', args=[c.code])),
            c.certificate_date.strftime('%d/%m/%Y'),
            timezone.localtime(c.issued_at).strftime('%d/%m/%Y %H:%M'),
            timezone.localtime(c.emailed_at).strftime('%d/%m/%Y %H:%M') if c.emailed_at else '',
            'Revogado' if c.is_revoked else 'Válido',
            timezone.localtime(c.revoked_at).strftime('%d/%m/%Y %H:%M') if c.revoked_at else '',
            c.revoked_reason,
        ])  # fmt: skip
    return response


@login_required
@require_POST
def batch_email(request, pk):
    batch = _get_batch(request, pk)
    if batch is None:
        raise PermissionDenied(FORBIDDEN)
    if closed := _closed(request, batch):
        return closed
    if not EmailSettings.load().enabled:
        messages.error(request, 'O envio de e-mails está desligado. Configure em Configurações > E-mail.')
        return redirect('batch_detail', pk=batch.pk)
    queued, waiting = queue_emails(batch, resend=request.POST.get('resend') == '1', request=request)
    if queued:
        plural = 's' if queued != 1 else ''
        messages.success(
            request,
            f'{queued} certificado{plural} na fila de envio. Os e-mails saem em instantes, um por um; '
            'atualize esta página para acompanhar.',
        )
    elif waiting:
        messages.info(
            request,
            f'{waiting} e-mail{"s" if waiting != 1 else ""} deste lote já estava{"m" if waiting != 1 else ""} na fila.',
        )
    else:
        messages.info(request, 'Não havia certificados pendentes de envio.')
    return redirect('batch_detail', pk=batch.pk)


@login_required
def batch_template(request):
    """Planilha modelo, já com o cabeçalho que a importação reconhece. ?modelo=completo: com as colunas do evento."""
    if not can_manage_batch(request.user):
        raise PermissionDenied(FORBIDDEN)
    from openpyxl import Workbook
    from openpyxl.styles import Font

    full = request.GET.get('modelo') == 'completo'
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = 'Participantes'
    event_headers = list(EVENT_COLUMNS.values()) if full else []
    headers = event_headers + ['Nome', 'Matrícula', 'Turma', 'E-mail']
    sheet.append(headers)
    widths = [36, 10, 13, 13, 14, 20][: len(event_headers)] + [40, 16, 18, 36]
    for index, (cell, width) in enumerate(zip(sheet[1], widths, strict=True), start=1):
        cell.font = Font(bold=True)
        sheet.column_dimensions[cell.column_letter].width = width
        for row in range(2, 501):
            if headers[index - 1] == 'Matrícula':
                sheet.cell(row=row, column=index).number_format = '@'  # texto: não perde zeros à esquerda
            elif headers[index - 1] in ('Início', 'Término', 'Data do certificado'):
                sheet.cell(row=row, column=index).number_format = 'DD/MM/YYYY'
    if full:
        guide = workbook.create_sheet('Como preencher')
        guide.column_dimensions['A'].width = 24
        guide.column_dimensions['B'].width = 90
        for row in [
            ('Coluna', 'Como preencher'),
            ('Evento', 'Obrigatória. Título como sai no certificado, em negrito e entre aspas.'),
            ('Tipo', 'Curso, Evento ou Agenda. Em branco: Evento.'),
            ('Início', 'Obrigatória. Data no formato dd/mm/aaaa.'),
            ('Término', 'Em branco: o mesmo dia do início.'),
            ('Carga horária', 'Só o número de horas (ex.: 20). Em branco: o certificado não cita a duração.'),
            ('Data do certificado', 'A data impressa no certificado. Em branco: a data de término.'),
            ('Nome', 'Obrigatória. Um participante por linha.'),
            ('Matrícula, Turma, E-mail', 'Opcionais. Com e-mail, dá para enviar o certificado por e-mail.'),
            ('', ''),
            ('Um evento por planilha', 'As colunas do evento podem ser preenchidas só na primeira linha. Se uma '
             'linha trouxer um evento diferente, a planilha é recusada.'),
            ('Esta aba', 'É só de instruções e não é lida. Os participantes vão na aba "Participantes".'),
        ]:  # fmt: skip
            guide.append(row)
        for cell in guide[1]:
            cell.font = Font(bold=True)
    buffer = io.BytesIO()
    workbook.save(buffer)
    response = HttpResponse(
        buffer.getvalue(), content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    )
    filename = 'modelo-planilha-completa.xlsx' if full else 'modelo-participantes.xlsx'
    response['Content-Disposition'] = f'attachment; filename="{filename}"'
    return response
