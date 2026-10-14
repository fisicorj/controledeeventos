"""Leitura da planilha de participantes da emissão avulsa (.xlsx ou .csv).

As colunas são reconhecidas pelo nome do cabeçalho, com folga: "Nome (conforme base de acesso)",
"Matrícula", "RA", "Turma", "E-mail"... A ordem não importa e colunas desconhecidas são ignoradas.
"""

import csv
import io
import re
import unicodedata
from dataclasses import asdict, dataclass, field
from datetime import date, datetime

from django.core.exceptions import ValidationError
from django.core.validators import validate_email

MAX_ROWS = 3000
MAX_FILE_SIZE = 5 * 1024 * 1024
HEADER_SCAN_ROWS = 10

# Conectivos que ficam em minúsculas no meio do nome
LOWER_WORDS = {'da', 'das', 'de', 'do', 'dos', 'e', 'di', 'du', 'del', 'della', 'van', 'von', 'y'}


class SpreadsheetError(Exception):
    """Problema que impede a leitura da planilha inteira (arquivo inválido, sem coluna de nome...)."""


@dataclass
class Row:
    line: int
    name: str
    original_name: str
    document: str = ''
    group: str = ''
    email: str = ''
    errors: list = field(default_factory=list)
    warnings: list = field(default_factory=list)

    @property
    def ok(self):
        return not self.errors

    def to_dict(self):
        return asdict(self)


def _plain(text):
    text = unicodedata.normalize('NFKD', str(text)).encode('ascii', 'ignore').decode()
    return re.sub(r'[^a-z0-9]+', ' ', text.lower()).strip()


START_HEADERS = {
    'data',
    'inicio',
    'data de inicio',
    'data inicial',
    'data do inicio',
    'data do evento',
    'inicio do evento',
}
END_HEADERS = {'termino', 'fim', 'data de termino', 'data final', 'data fim', 'data do termino', 'data do fim',
               'termino do evento', 'fim do evento'}  # fmt: skip


# Regras de reconhecimento das colunas, na ordem em que são testadas: a primeira que bater decide.
# Cada regra recebe o cabeçalho sem acento e em minúsculas (h) e as palavras dele (w).
COLUMN_RULES = [
    ('email', lambda h, w: 'mail' in h),
    ('group', lambda h, w: w[0] in ('turma', 'classe') or 'turma' in w),
    # Colunas do evento (planilha completa). Vêm antes de "nome" por causa de "Nome do evento".
    ('issue_date', lambda h, w: 'certificado' in w or 'emissao' in w),
    ('workload_hours', lambda h, w: 'carga' in w or h in ('horas', 'duracao', 'duracao horas')),
    ('start_date', lambda h, w: h in START_HEADERS),
    ('end_date', lambda h, w: h in END_HEADERS),
    ('activity_type', lambda h, w: w[0] == 'tipo'),
    ('title', lambda h, w: 'evento' in w or w[0] == 'titulo' or h == 'atividade'),
    ('name', lambda h, w: w[0] == 'nome' or h in ('participante', 'aluno', 'aluna', 'estudante')),
    ('document', lambda h, w: w[0] in ('matricula', 'ra', 'rm', 'prontuario') or h.startswith('registro academico')),
]


def column_kind(header):
    """Identifica a coluna pelo texto do cabeçalho (ver COLUMN_RULES)."""
    h = _plain(header or '')
    if not h:
        return None
    words = h.split()
    return next((kind for kind, rule in COLUMN_RULES if rule(h, words)), None)


def cell_text(value):
    """Converte o valor da célula em texto: 12546456.0 vira "12546456", espaços repetidos somem."""
    if value is None:
        return ''
    if isinstance(value, bool):
        return ''
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    return re.sub(r'\s+', ' ', str(value)).strip()


def smart_title(name):
    """Ajusta maiúsculas: "joaninha melo rego" -> "Joaninha Melo Rego", "MARIA DA SILVA" -> "Maria da Silva".
    Palavras com maiúsculas no meio (McDonald, DiCaprio) ficam como estão."""
    words = name.split(' ')
    out = []
    for i, word in enumerate(words):
        if not word:
            continue
        if not (word.islower() or word.isupper()):
            out.append(word)
            continue
        low = word.lower()
        if i > 0 and low in LOWER_WORDS:
            out.append(low)
        else:
            out.append('-'.join(part[:1].upper() + part[1:] for part in low.split('-')))
    return ' '.join(out)


def _read_xlsx(data):
    from openpyxl import load_workbook

    try:
        workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except Exception as exc:
        raise SpreadsheetError('Não foi possível abrir a planilha. Confira se o arquivo é um .xlsx válido.') from exc
    try:
        sheet = workbook.worksheets[0]
        rows = []
        for row in sheet.iter_rows(values_only=True):
            rows.append(list(row))
            if len(rows) > MAX_ROWS + HEADER_SCAN_ROWS + 1:
                break
        return rows
    finally:
        workbook.close()


def _read_csv(data):
    for encoding in ('utf-8-sig', 'cp1252'):
        try:
            text = data.decode(encoding)
            break
        except UnicodeDecodeError:
            continue
    else:
        raise SpreadsheetError('Não foi possível ler o CSV. Salve o arquivo em UTF-8.')
    first_line = text.split('\n', 1)[0]
    delimiter = max(';,\t', key=first_line.count)
    return list(csv.reader(io.StringIO(text), delimiter=delimiter))


def read_table(uploaded):
    """Devolve as linhas do arquivo enviado como listas de valores."""
    if uploaded.size > MAX_FILE_SIZE:
        raise SpreadsheetError('O arquivo passa de 5 MB.')
    name = uploaded.name.lower()
    data = uploaded.read()
    if name.endswith('.xlsx'):
        return _read_xlsx(data)
    if name.endswith('.csv'):
        return _read_csv(data)
    if name.endswith('.xls'):
        raise SpreadsheetError('Arquivos .xls (Excel antigo) não são aceitos. No Excel, use "Salvar como" > .xlsx.')
    raise SpreadsheetError('Envie a planilha em .xlsx ou .csv.')


def _find_header(table):
    for index, row in enumerate(table[:HEADER_SCAN_ROWS]):
        columns = {}
        for col, value in enumerate(row):
            kind = column_kind(cell_text(value))
            if kind and kind not in columns:
                columns[kind] = col
        if 'name' in columns:
            return index, columns
    raise SpreadsheetError(
        'Não encontrei a coluna com o nome dos participantes. A primeira linha deve ter um cabeçalho '
        'como "Nome", "Matrícula", "Turma" e, se quiser enviar por e-mail, "E-mail".'
    )


def _check_fields(row):
    """Tamanhos e formato do e-mail."""
    if not row.name:
        row.errors.append('Sem nome')
    elif len(row.name) > 300:
        row.errors.append('Nome com mais de 300 caracteres')
    if len(row.document) > 50:
        row.errors.append('Matrícula com mais de 50 caracteres')
    if len(row.group) > 100:
        row.group = row.group[:100]
    if row.email:
        try:
            validate_email(row.email)
        except ValidationError:
            row.errors.append('E-mail inválido')


def _check_document(row, has_document_column):
    """Matrícula sem nenhum dígito ("Externo", "Convidado", "-", "N/A"...) não é matrícula: não vai para o
    certificado e não serve para achar repetidos (senão todos os externos virariam a mesma pessoa)."""
    if row.document and not any(ch.isdigit() for ch in row.document):
        if row.name:
            row.warnings.append(f'Sem matrícula ("{row.document}")')
        row.document = ''
    elif has_document_column and not row.document and row.name:
        row.warnings.append('Sem matrícula')


def _check_duplicate(row, existing, seen):
    """Já emitido neste lote, ou repetido dentro da própria planilha (seen: chave -> linha)."""
    key = participant_key(row.name, row.document)
    if row.name and key in existing:
        row.errors.append('Já tem certificado neste lote')
    elif row.name and key in seen:
        row.errors.append(f'Repetido (linha {seen[key]})')
    elif row.name:
        seen[key] = row.line


def parse_participants(table, fix_case=True, existing=None):
    """Valida as linhas da planilha.

    existing: participantes já emitidos no lote, como conjunto de chaves (ver participant_key).
    Devolve (linhas, colunas encontradas).
    """
    existing = existing or set()
    header_index, columns = _find_header(table)

    def get(row, kind):
        col = columns.get(kind)
        return cell_text(row[col]) if col is not None and col < len(row) else ''

    rows, seen = [], {}
    for offset, raw in enumerate(table[header_index + 1 :], start=header_index + 2):
        values = [cell_text(v) for v in raw]
        if not any(values):
            continue
        original = get(raw, 'name')
        row = Row(
            line=offset,
            name=smart_title(original) if fix_case else original,
            original_name=original,
            document=get(raw, 'document'),
            group=get(raw, 'group'),
            email=get(raw, 'email').lower(),
        )
        _check_fields(row)
        _check_document(row, has_document_column='document' in columns)
        _check_duplicate(row, existing, seen)
        rows.append(row)
        if len(rows) > MAX_ROWS:
            raise SpreadsheetError(f'A planilha passa de {MAX_ROWS} participantes. Divida em mais de um arquivo.')
    if not rows:
        raise SpreadsheetError('A planilha não tem participantes abaixo do cabeçalho.')
    return rows, set(columns)


# ---------- Planilha completa: os dados do evento vêm nas colunas da própria planilha

EVENT_COLUMNS = {
    'title': 'Evento',
    'activity_type': 'Tipo',
    'start_date': 'Início',
    'end_date': 'Término',
    'workload_hours': 'Carga horária',
    'issue_date': 'Data do certificado',
}
TYPE_NAMES = {'curso': 'COURSE', 'course': 'COURSE', 'evento': 'EVENT', 'event': 'EVENT', 'agenda': 'APPOINTMENT',
              'appointment': 'APPOINTMENT'}  # fmt: skip
DATE_FORMATS = ('%d/%m/%Y', '%d/%m/%y', '%Y-%m-%d', '%d-%m-%Y', '%d.%m.%Y')


def parse_date(value):
    """Data da célula: data do Excel ou texto dd/mm/aaaa (também aaaa-mm-dd). Devolve date, None (vazia) ou
    levanta ValueError."""
    if value is None or value == '':
        return None
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    text = cell_text(value).split(' ')[0]
    if not text:
        return None
    for fmt in DATE_FORMATS:
        try:
            return datetime.strptime(text, fmt).date()
        except ValueError:
            continue
    raise ValueError(text)


def _event_value(kind, raw):
    """Valor da célula convertido para o campo do evento (None se a célula está vazia)."""
    if kind in ('start_date', 'end_date', 'issue_date'):
        return parse_date(raw)
    text = cell_text(raw)
    if not text:
        return None
    if kind == 'workload_hours':
        digits = re.match(r'^(\d+)(?:[.,]0+)?(?:\s*h(?:oras?)?)?$', text.lower())
        if not digits:
            raise ValueError(text)
        return int(digits.group(1))
    if kind == 'activity_type':
        code = TYPE_NAMES.get(_plain(text))
        if code is None:
            raise ValueError(text)
        return code
    return text


def _show(value):
    return value.strftime('%d/%m/%Y') if isinstance(value, date) else str(value)


def read_event(table):
    """Dados do evento de uma planilha completa (um evento por planilha).

    As colunas do evento podem vir preenchidas em todas as linhas ou só na primeira: em branco vale o que veio
    antes. Se alguma linha trouxer um valor diferente, a planilha é recusada (é outro evento, ou erro de
    digitação), para ninguém receber certificado com dados trocados. Devolve um dicionário com os campos do
    formulário do evento."""
    header_index, columns = _find_header(table)
    missing = [EVENT_COLUMNS[k] for k in ('title', 'start_date') if k not in columns]
    if missing:
        raise SpreadsheetError(
            f'Faltam as colunas do evento: {", ".join(missing)}. Na planilha completa, o cabeçalho precisa ter '
            'Evento, Início e Nome (as outras colunas do modelo são opcionais).'
        )
    found = {}  # campo: (valor, linha)
    for line, raw in enumerate(table[header_index + 1 :], start=header_index + 2):
        if not any(cell_text(v) for v in raw):
            continue
        for kind, label in EVENT_COLUMNS.items():
            col = columns.get(kind)
            if col is None or col >= len(raw):
                continue
            try:
                value = _event_value(kind, raw[col])
            except ValueError as exc:
                hint = {
                    'workload_hours': 'Use só o número de horas, ex.: 20.',
                    'activity_type': 'Use Curso, Evento ou Agenda.',
                }.get(kind, 'Use o formato dd/mm/aaaa.')
                raise SpreadsheetError(f'Linha {line}, coluna {label}: "{exc}" não foi entendido. {hint}') from exc
            if value is None:
                continue
            if kind not in found:
                found[kind] = (value, line)
            elif found[kind][0] != value:
                first, first_line = found[kind]
                raise SpreadsheetError(
                    f'A linha {line} traz {label} "{_show(value)}", '
                    f'diferente da linha {first_line} ("{_show(first)}"). '
                    'A planilha completa é de um evento só: deixe a coluna igual em todas as linhas (ou preencha só '
                    'a primeira) e use outra planilha para outro evento.'
                )
    if 'title' not in found or 'start_date' not in found:
        absent = [EVENT_COLUMNS[k] for k in ('title', 'start_date') if k not in found]
        raise SpreadsheetError(f'Preencha {" e ".join(absent)} pelo menos na primeira linha de participante.')
    event = {kind: value for kind, (value, _line) in found.items()}
    event.setdefault('activity_type', 'EVENT')
    event.setdefault('end_date', event['start_date'])
    event.setdefault('issue_date', event['end_date'])
    event.setdefault('workload_hours', None)
    return event


def participant_key(name, document):
    """Identifica o participante dentro do lote: pela matrícula, ou pelo nome quando não há matrícula."""
    return ('doc', document.strip().upper()) if document.strip() else ('name', _plain(name))
